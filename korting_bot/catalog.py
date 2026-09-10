# -*- coding: utf-8 -*-
"""Собрать локальный индекс из выгрузки Bitrix на Яндекс.Диске."""
from __future__ import annotations

import html
import json
import re
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from .paths import CATALOG_PATH, DATA_DIR, FEED_URL, SITE_BASE, YML_ENRICH_URL

USER_AGENT = "KortingSpecsBot/1.0"
DISK_RESOURCES = "https://cloud-api.yandex.net/v1/disk/public/resources"
DISK_DOWNLOAD = "https://cloud-api.yandex.net/v1/disk/public/resources/download"

SHEET_CATEGORY = {
    "Духовые шкафы": "Духовые шкафы",
    "Варочные поверхности": "Варочные поверхности",
    "Вытяжки": "Вытяжки",
    "Посудомоечные машины": "Посудомоечные машины",
    "Стиральные машины": "Стиральные машины",
    "Холодильники": "Холодильники",
    "Микроволновки": "Микроволновые печи",
    "Сушилки": "Сушильные машины",
    "Морозильные камеры": "Морозильные камеры",
    "Кофемашины": "Кофемашины",
    "Винные шкафы": "Винные шкафы",
    "Ящики для подогрева": "Ящики для подогрева",
    "Аксессуары": "Аксессуары",
    "Кофеварки": "Кофеварки",
    "Чайники": "Чайники",
    "Погружные блендеры": "Погружные блендеры",
    "Настольные блендеры": "Настольные блендеры",
    "Мясорубки": "Мясорубки",
    "Мультивыпечки": "Мультивыпечки",
    "Мельницы и штопоры": "Мельницы и штопоры",
    "Кухонные машины": "Кухонные машины",
    "Кухонные комбайны": "Кухонные комбайны",
    "Кофемолки": "Кофемолки",
    "Кофейные станции": "Кофейные станции",
    "Грили": "Грили",
    "Весы напольные": "Весы напольные",
    "Весы кухонные": "Весы кухонные",
    "Стир-суш колонны": "Стирально-сушильные колонны",
}

SKIP_PARAM_NAMES = {
    "id элемента",
    "ean",
    "код из 1с",
    "модель",
    "наименование элемента",
    "url страницы детального просмотра",
    "внешний код",
    "в архиве",
    "описание для анонса",
    "детальное описание",
    "новинка",
    "популярный",
    "представлено в шоу-руме",
    "спецпредложение",
    "оставлять новинкой",
    "выставочный образец",
    "причины уценки",
    "запрет на обновление из 1с",
    "фотогаллерея",
    "фотогалерея",
    "выгружать в директ",
    "эксклюзив",
    "доступен к покупке",
    "покупка или подписка",
    "поисковый индекс",
    "номер фото для директа",
    "видео",
    "сортировка видео",
    "инструкция",
    "общие рекомендации по установке",
    "схема встраивания",
    "сочетание с другими приборами",
    "не забудьте купить",
    "id поста блога для комментариев",
    "количество комментариев",
    "цена для фильтра",
    "графический рич контент",
    "дата окончания новинки",
    "не выводить ссылку на интернет магазин",
    "не выводить на сайте магазина",
    "не выгружать в yandex market",
    "описание | видео",
    "привязка к цветам",
    "фото | преимущества и особенности",
    "фото для директа",
    "фотогалерея webp",
    "основные характеристики",
    "дополнительные характеристики",
    "наличие",
}
SKIP_PARAM_PREFIXES = (
    "фото",
    "видео",
    "инструкция",
    "схема",
    "вывод на главной",
    "ссылка на товар",
    "не выводить",
    "не выгружать",
    "номер фото",
)


def _writable_dir() -> Path:
    for folder in (DATA_DIR, Path("/tmp/korting"), Path(tempfile.gettempdir()) / "korting"):
        try:
            folder.mkdir(parents=True, exist_ok=True)
            probe = folder / ".write_test"
            probe.write_text("ok", encoding="utf-8")
            probe.unlink()
            return folder
        except OSError:
            continue
    return Path(tempfile.gettempdir())


def _http_json(url: str) -> dict:
    req = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8"))


def _http_bytes(url: str) -> bytes:
    req = Request(url, headers={"User-Agent": USER_AGENT})
    with urlopen(req, timeout=180) as resp:
        return resp.read()


def _feed_date(meta: dict) -> str:
    raw = (meta.get("modified") or meta.get("created") or "").strip()
    if not raw:
        return ""
    return raw.replace("T", " ")[:19]


def _pick_xlsx_item(meta: dict) -> dict:
    if meta.get("type") == "file":
        return meta
    items = (meta.get("_embedded") or {}).get("items") or []
    xlsx = [it for it in items if str(it.get("name") or "").lower().endswith(".xlsx")]
    if not xlsx:
        raise RuntimeError("на Яндекс.Диске нет xlsx-файла")
    xlsx.sort(key=lambda it: it.get("modified") or "", reverse=True)
    return xlsx[0]


def download_feed(dest: Path, public_url: str = FEED_URL) -> tuple[Path, str]:
    dest.parent.mkdir(parents=True, exist_ok=True)
    q = urlencode({"public_key": public_url})
    meta = _http_json(f"{DISK_RESOURCES}?{q}&limit=100")
    item = _pick_xlsx_item(meta)
    if item.get("type") == "file" and item.get("path") in {"/", "", None}:
        href = _http_json(f"{DISK_DOWNLOAD}?{q}")["href"]
    else:
        path = item.get("path") or f"/{item.get('name')}"
        href = _http_json(f"{DISK_DOWNLOAD}?{urlencode({'public_key': public_url, 'path': path})}")["href"]
    dest.write_bytes(_http_bytes(href))
    return dest, _feed_date(item)


def download_yml(dest: Path, url: str = FEED_URL) -> Path:
    path, _date = download_feed(dest, url)
    return path


def _clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Y" if value else "N"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).rstrip("0").rstrip(".")
    text = html.unescape(str(value).replace("\xa0", " ")).replace("&nbsp;", " ")
    return re.sub(r"\s+", " ", text).strip()


def _header_name(raw: str) -> str:
    s = _clean_text(raw)
    s = re.sub(r"\s*\{[^}]+\}\s*$", "", s)
    s = re.sub(r"\s*\((?:описание|description)\)\s*$", "", s, flags=re.I)
    s = re.sub(r"\s*\[[^\]]+\]\s*$", "", s)
    return _clean_text(s)


def _norm_name(name: str) -> str:
    return _clean_text(name).lower().replace("ё", "е")


def _is_enum_ids(val: str) -> bool:
    parts = [p for p in (val or "").split("|") if p]
    return len(parts) >= 2 and all(p.isdigit() for p in parts)


def _skip_param(name: str) -> bool:
    n = _norm_name(name)
    if not n or n in SKIP_PARAM_NAMES:
        return True
    return any(n.startswith(p) for p in SKIP_PARAM_PREFIXES)


def _product_url(raw: str) -> str:
    url = _clean_text(raw).split("?")[0]
    if not url:
        return ""
    if url.startswith("http://") or url.startswith("https://"):
        return url
    if url.startswith("/"):
        return SITE_BASE + url
    return SITE_BASE + "/" + url.lstrip("/")


def parse_xlsx(path: Path, source: str = FEED_URL, feed_date: str = "") -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    products = []
    seen: set[str] = set()
    try:
        for sheet_name in wb.sheetnames:
            category = SHEET_CATEGORY.get(sheet_name, sheet_name)
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            header_row = next(rows, None)
            if not header_row:
                continue
            headers = [_header_name(h) for h in header_row]
            for row in rows:
                values: dict[str, str] = {}
                for i, key in enumerate(headers):
                    if not key or i >= len(row):
                        continue
                    val = _clean_text(row[i])
                    if not val or key in values:
                        continue
                    values[key] = val
                model = values.get("Модель") or ""
                if not model:
                    continue
                pid = values.get("ID элемента") or model
                uniq = f"{pid}:{model}"
                if uniq in seen:
                    continue
                seen.add(uniq)
                params = {}
                for key, val in values.items():
                    if _skip_param(key) or _is_enum_ids(val):
                        continue
                    params[key] = val
                products.append({
                    "id": pid,
                    "model": model,
                    "name": values.get("Наименование элемента") or model,
                    "category_id": sheet_name,
                    "category": category,
                    "url": _product_url(values.get("URL страницы детального просмотра") or ""),
                    "params": params,
                })
    finally:
        wb.close()
    return {"source": source, "feed_date": feed_date, "products": products}


def _full_category(cats: dict, cid: str) -> str:
    info = cats.get(cid) or {}
    name = info.get("name") or cid
    parent = info.get("parent")
    if parent and cats.get(parent, {}).get("name"):
        return cats[parent]["name"] + " / " + name
    return name


def parse_yml(path: Path) -> dict:
    tree = ET.parse(path)
    shop = tree.getroot().find("shop")
    feed_date = (tree.getroot().get("date") or "").strip()
    cats = {
        c.get("id"): {"name": (c.text or "").strip(), "parent": c.get("parentId") or ""}
        for c in shop.find("categories")
    }
    products = []
    for o in shop.find("offers"):
        cid = (o.findtext("categoryId") or "").strip()
        params = {}
        for p in o.findall("param"):
            n = (p.get("name") or "").strip()
            if not n:
                continue
            params[n] = (p.text or "").strip()
        model = (o.findtext("model") or "").strip()
        products.append({
            "id": o.get("id") or "",
            "model": model,
            "name": (o.findtext("name") or "").strip(),
            "category_id": cid,
            "category": _full_category(cats, cid),
            "url": (o.findtext("url") or "").split("?")[0],
            "params": params,
        })
    return {"source": str(path), "feed_date": feed_date, "products": products}


def parse_feed_file(path: Path, source: str = FEED_URL, feed_date: str = "") -> dict:
    suf = path.suffix.lower()
    if suf in {".xlsx", ".xlsm"}:
        return parse_xlsx(path, source=source, feed_date=feed_date)
    return parse_yml(path)


def save_catalog(data: dict, path: Path = CATALOG_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return path


def refresh_from_file(feed_path: Path, catalog_path: Path = CATALOG_PATH) -> dict:
    data = parse_feed_file(feed_path)
    save_catalog(data, catalog_path)
    return data


def _enrich_from_yml(data: dict) -> dict:
    """Подставить текстовые преимущества из YML, где в Excel только ID справочника."""
    folder = _writable_dir()
    path = folder / "enrich.yml"
    try:
        req = Request(YML_ENRICH_URL, headers={"User-Agent": USER_AGENT})
        with urlopen(req, timeout=180) as resp:
            path.write_bytes(resp.read())
        yml = parse_yml(path)
    except Exception as exc:
        print(f"yml enrich skipped: {exc}", flush=True)
        return data
    finally:
        try:
            path.unlink()
        except OSError:
            pass
    by_model = {
        (p.get("model") or "").strip(): p
        for p in yml.get("products") or []
        if (p.get("model") or "").strip()
    }
    filled = 0
    for p in data["products"]:
        src = by_model.get((p.get("model") or "").strip())
        if not src:
            continue
        params = p.setdefault("params", {})
        for key, val in (src.get("params") or {}).items():
            text = (val or "").strip()
            if not text or _skip_param(key) or _is_enum_ids(text):
                continue
            cur = params.get(key)
            if cur is None or _is_enum_ids(cur):
                params[key] = text
                filled += 1
        if not p.get("url") and src.get("url"):
            p["url"] = src["url"].split("?")[0]
    print(f"yml enrich fields {filled}", flush=True)
    return data


def refresh_from_url(catalog_path: Path | None = None, yml_path: Path | None = None) -> dict:
    folder = _writable_dir()
    catalog_path = catalog_path or (folder / "catalog.json")
    feed_path = yml_path or (folder / "feed.xlsx")
    downloaded, feed_date = download_feed(feed_path)
    data = parse_xlsx(downloaded, source=FEED_URL, feed_date=feed_date)
    data = _enrich_from_yml(data)
    save_catalog(data, catalog_path)
    try:
        downloaded.unlink()
    except OSError:
        pass
    return data


def ensure_catalog(path: Path = CATALOG_PATH) -> Path:
    if path.is_file() and path.stat().st_size > 0:
        return path
    dest = path
    try:
        dest.parent.mkdir(parents=True, exist_ok=True)
        probe = dest.parent / ".write_test"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink()
    except OSError:
        dest = _writable_dir() / "catalog.json"
        if dest.is_file() and dest.stat().st_size > 0:
            return dest
    print(f"catalog missing at {path}, downloading feed to {dest}", flush=True)
    refresh_from_url(dest)
    return dest


def load_catalog(path: Path = CATALOG_PATH) -> dict:
    resolved = ensure_catalog(path)
    return json.loads(resolved.read_text(encoding="utf-8"))
