# -*- coding: utf-8 -*-
"""Build the public USP product catalog used by GitHub Pages."""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from difflib import SequenceMatcher
import json
from pathlib import Path
import re
import tempfile
from typing import Any

from openpyxl import load_workbook

from korting_bot.catalog import SITE_BASE, download_feed


ROOT = Path(__file__).resolve().parent
DEFAULT_GLOSSARY = (
    ROOT.parent
    / "Финальные таблицы"
    / "Сборка тексты и USP"
    / "FULL_GLOSSARY_USP_TEXTS.xlsx"
)
OUTPUT_DIR = ROOT / "docs" / "USP_CATALOG"

CATEGORY_SHEETS = [
    ("Варочные поверхности — газовые", "GLOSS_ВАРОЧНЫЕ ГАЗОВЫЕ", "USP_MODELS_ВАРОЧНЫЕ ГАЗ"),
    (
        "Варочные поверхности — электрические",
        "GLOSS_ВАРОЧНЫЕ ЭЛЕКТРИЧЕСКИЕ",
        "USP_MODELS_ВАРОЧНЫЕ ЭЛ",
    ),
    ("Духовые шкафы", "GLOSS_ДУХОВЫЕ ШКАФЫ", "USP_MODELS_ДУХОВЫЕ ШКАФЫ"),
    (
        "Стиральные машины — встраиваемые",
        "GLOSS_СТИР. МАШИНЫ (ВСТР)",
        "USP_MODELS_СТИР.МАШ.(ВСТР)",
    ),
    (
        "Стиральные машины — отдельностоящие",
        "GLOSS_СТИР. МАШИНЫ (ОТД)",
        "USP_MODELS_СТИР.МАШ.(ОТД)",
    ),
    (
        "Холодильники — встраиваемые",
        "GLOSS_ХОЛОДИЛЬНИКИ (ВСТР)",
        "USP_MODELS_ХОЛОДИЛЬН.(ВСТР)",
    ),
    (
        "Холодильники — отдельностоящие",
        "GLOSS_ХОЛОДИЛЬНИКИ (ОТД)",
        "USP_MODELS_ХОЛОДИЛЬН.(ОТД)",
    ),
    ("Сушильные машины", "GLOSS_СУШИЛЬНЫЕ МАШИНЫ", "USP_MODELS_СУШИЛЬНЫЕ"),
    (
        "Стирально-сушильные колонны",
        "GLOSS_СТИР.-СУШ. КОЛОННЫ",
        "USP_MODELS_СТИР.-СУШ.КОЛ.",
    ),
    (
        "Посудомоечные машины — встраиваемые",
        "GLOSS_ПОСУД. МАШИНЫ (ВСТР)",
        "USP_MODELS_ПОСУД.МАШ.(ВСТР)",
    ),
    (
        "Посудомоечные машины — отдельностоящие",
        "GLOSS_ПОСУД. МАШИНЫ (ОТД)",
        "USP_MODELS_ПОСУД.МАШ.(ОТД)",
    ),
    ("Вытяжки", "GLOSS_ВЫТЯЖКИ", "USP_MODELS_ВЫТЯЖКИ"),
    (
        "Морозильные камеры — встраиваемые",
        "GLOSS_МОРОЗ. КАМЕРЫ (ВСТР)",
        "USP_MODELS_МОРОЗ.КАМ.(ВСТР)",
    ),
    (
        "Морозильные камеры — отдельностоящие",
        "GLOSS_МОРОЗ. КАМЕРЫ (ОТД)",
        "USP_MODELS_МОРОЗ.КАМ.(ОТД)",
    ),
    (
        "Микроволновые печи — встраиваемые",
        "GLOSS_МИКРОВОЛН. ПЕЧИ (ВСТР)",
        "USP_MODELS_СВЧ(ВСТР)",
    ),
    (
        "Микроволновые печи — отдельностоящие",
        "GLOSS_МИКРОВОЛН. ПЕЧИ (МБТ)",
        "USP_MODELS_СВЧ(МБТ)",
    ),
    ("Винные шкафы", "GLOSS_ВИННЫЕ ШКАФЫ", "USP_MODELS_ВИННЫЕ ШКАФЫ"),
    (
        "Ящики для подогрева",
        "GLOSS_ЯЩИКИ ДЛЯ ПОДОГРЕВА",
        "USP_MODELS_ЯЩИКИ ПОДОГРЕВА",
    ),
    ("Малая бытовая техника", "GLOSS_МБТ", "USP_MODELS_МБТ"),
]

NOISE_WORDS = {
    "и",
    "с",
    "со",
    "в",
    "во",
    "для",
    "на",
    "по",
    "из",
    "до",
    "от",
    "а",
}
MODEL_PREFIX = re.compile(
    r"^(?:"
    r"встраиваем(?:ая|ый|ое)\s+|"
    r"отдельностоящ(?:ая|ий|ее)\s+|"
    r"стиральная\s+машина(?:\s+с\s+сушкой)?\s+|"
    r"стирально-сушильная\s+(?:машина|колонна)\s+|"
    r"сушильная\s+машина\s+|"
    r"посудомоечная\s+машина\s+|"
    r"микроволновая\s+печь\s+|"
    r"морозильная\s+камера\s+|"
    r"холодильник\s+|"
    r"духовой\s+шкаф\s+|"
    r"варочная\s+поверхность\s+|"
    r"вытяжка\s+|"
    r"винный\s+шкаф\s+|"
    r"ящик\s+для\s+подогрева\s+"
    r")+",
    re.IGNORECASE,
)


def clean(value: object) -> str:
    text = str(value or "").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize(value: object) -> str:
    text = clean(value).lower().replace("ё", "е")
    text = text.replace("led", "светодиод")
    text = re.sub(r"[«»\"'`´]", "", text)
    text = re.sub(r"[^0-9a-zа-я%+]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def tokens(value: object) -> set[str]:
    return {
        token
        for token in normalize(value).split()
        if token not in NOISE_WORDS and len(token) > 1
    }


def canonical_model(value: object) -> str:
    text = clean(value)
    previous = None
    while text != previous:
        previous = text
        text = MODEL_PREFIX.sub("", text).strip()
    return re.sub(r"\s+", " ", text.upper().replace("Ё", "Е"))


def public_url(value: object) -> str:
    raw = clean(value).split("|")[0]
    if not raw:
        return ""
    if raw.startswith(("http://", "https://")):
        return raw
    return SITE_BASE + "/" + raw.lstrip("/")


def locate_header(headers: list[object], *needles: str) -> int | None:
    upper_needles = tuple(needle.upper() for needle in needles)
    for index, value in enumerate(headers):
        raw = str(value or "").upper()
        if any(needle in raw for needle in upper_needles):
            return index
    return None


def load_feed_products(feed_path: Path) -> dict[str, dict[str, str]]:
    # Normal mode is much faster here than repeated access to read-only rows
    # in a wide Bitrix export.
    workbook = load_workbook(feed_path, read_only=False, data_only=True)
    products: dict[str, dict[str, str]] = {}
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, ()) or ())
            if not headers:
                continue
            model_index = 3
            photo_index = locate_header(headers, "[DIRECT_PHOTO]")
            url_index = locate_header(headers, "[IE_DETAIL_PAGE_URL]", "URL СТРАНИЦЫ")
            name_index = locate_header(headers, "[IE_NAME]", "НАИМЕНОВАНИЕ ЭЛЕМЕНТА")
            if photo_index is None:
                continue
            for row in rows:
                if len(row) <= model_index:
                    continue
                model = canonical_model(row[model_index])
                if not model:
                    continue
                photo = public_url(row[photo_index] if len(row) > photo_index else "")
                url = public_url(row[url_index] if url_index is not None and len(row) > url_index else "")
                name = clean(row[name_index] if name_index is not None and len(row) > name_index else "")
                candidate = {
                    "model": clean(row[model_index]),
                    "name": name,
                    "photo": photo,
                    "url": url,
                    "feed_category": sheet.title,
                }
                current = products.get(model)
                if current is None or (
                    bool(candidate["photo"]),
                    bool(candidate["url"]),
                ) > (
                    bool(current["photo"]),
                    bool(current["url"]),
                ):
                    products[model] = candidate
    finally:
        workbook.close()
    return products


def load_glossary(sheet) -> list[dict[str, str]]:
    headers = {
        clean(sheet.cell(1, col).value): col
        for col in range(1, sheet.max_column + 1)
    }
    name_col = headers.get("USP NAME", 5)
    description_col = headers.get("USP DESCRIPTION", 6)
    short_col = headers.get("USP DESCRIPTION_SHORT", 7)
    id_col = headers.get("ID GIVEN", 4)
    entries = []
    for row in range(2, sheet.max_row + 1):
        title = clean(sheet.cell(row, name_col).value)
        if not title:
            continue
        description = clean(sheet.cell(row, description_col).value)
        short = clean(sheet.cell(row, short_col).value)
        entries.append(
            {
                "id": clean(sheet.cell(row, id_col).value),
                "title": title,
                "description": description or short,
                "normalized": normalize(title),
            }
        )
    return entries


def best_glossary_match(title: str, entries: list[dict[str, str]]) -> tuple[dict[str, str] | None, str]:
    wanted = normalize(title)
    if not wanted:
        return None, "none"
    exact = [entry for entry in entries if entry["normalized"] == wanted]
    if exact:
        exact.sort(key=lambda item: bool(item["description"]), reverse=True)
        return exact[0], "exact"

    wanted_tokens = tokens(title)
    candidates: list[tuple[float, dict[str, str]]] = []
    for entry in entries:
        candidate = entry["normalized"]
        candidate_tokens = tokens(candidate)
        if not candidate_tokens or not wanted_tokens:
            continue
        overlap = len(wanted_tokens & candidate_tokens) / len(wanted_tokens | candidate_tokens)
        ratio = SequenceMatcher(None, wanted, candidate).ratio()
        contains = wanted in candidate or candidate in wanted
        score = ratio * 0.6 + overlap * 0.4 + (0.08 if contains else 0)
        if (contains and min(len(wanted), len(candidate)) >= 9 and score >= 0.65) or (
            ratio >= 0.82 and overlap >= 0.45
        ):
            candidates.append((score, entry))
    if not candidates:
        return None, "none"
    candidates.sort(key=lambda item: (item[0], bool(item[1]["description"])), reverse=True)
    best_score, best = candidates[0]
    if len(candidates) > 1 and best_score - candidates[1][0] < 0.025:
        return None, "ambiguous"
    return best, "fuzzy"


def build_products(
    workbook,
    feed_products: dict[str, dict[str, str]],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    result: list[dict[str, Any]] = []
    report = {
        "models": 0,
        "usps": 0,
        "described_usps": 0,
        "missing_photos": [],
        "missing_feed_models": [],
        "unmatched_usps": [],
        "ambiguous_usps": [],
        "fuzzy_matches": 0,
    }
    for order, (category, glossary_sheet, model_sheet) in enumerate(CATEGORY_SHEETS):
        glossary = load_glossary(workbook[glossary_sheet])
        sheet = workbook[model_sheet]
        for col in range(4, sheet.max_column + 1):
            model = clean(sheet.cell(1, col).value)
            if not model:
                continue
            model_key = canonical_model(model)
            feed = feed_products.get(model_key, {})
            usp_items = []
            seen = set()
            for row in range(3, sheet.max_row + 1, 2):
                title = clean(sheet.cell(row, col).value)
                if not title:
                    continue
                title_key = normalize(title)
                if not title_key or title_key in seen:
                    continue
                seen.add(title_key)
                direct_description = clean(sheet.cell(row + 1, col).value)
                matched, match_type = best_glossary_match(title, glossary)
                description = direct_description
                glossary_title = ""
                glossary_id = ""
                if matched:
                    glossary_title = matched["title"]
                    glossary_id = matched["id"]
                    if not description:
                        description = matched["description"]
                if match_type == "fuzzy":
                    report["fuzzy_matches"] += 1
                elif match_type == "ambiguous":
                    report["ambiguous_usps"].append(
                        {"category": category, "model": model, "usp": title}
                    )
                elif not matched:
                    report["unmatched_usps"].append(
                        {"category": category, "model": model, "usp": title}
                    )
                usp_items.append(
                    {
                        "title": title,
                        "description": description,
                        "glossary_title": glossary_title,
                        "id": glossary_id,
                    }
                )
                report["usps"] += 1
                if description:
                    report["described_usps"] += 1

            if not usp_items:
                continue
            if not feed:
                report["missing_feed_models"].append(
                    {"category": category, "model": model}
                )
            elif not feed.get("photo"):
                report["missing_photos"].append(
                    {"category": category, "model": model}
                )
            result.append(
                {
                    "category": category,
                    "category_order": order,
                    "model": model,
                    "name": feed.get("name") or model,
                    "photo": feed.get("photo") or "",
                    "url": feed.get("url") or "",
                    "usps": usp_items,
                }
            )
            report["models"] += 1

    result.sort(
        key=lambda item: (
            item["category_order"],
            canonical_model(item["model"]),
        )
    )
    return result, report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--glossary", type=Path, default=DEFAULT_GLOSSARY)
    parser.add_argument("--feed", type=Path)
    parser.add_argument("--output", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()

    if not args.glossary.is_file():
        raise FileNotFoundError(f"Glossary not found: {args.glossary}")

    feed_path = args.feed
    temporary_feed = False
    feed_date = ""
    if feed_path is None:
        feed_path = Path(tempfile.gettempdir()) / "korting_usp_feed.xlsx"
        feed_path, feed_date = download_feed(feed_path)
        temporary_feed = True

    feed_products = load_feed_products(feed_path)
    # Model sheets are wide and accessed by column. Normal mode avoids the
    # expensive XML rescans caused by random cell access in read-only mode.
    workbook = load_workbook(args.glossary, read_only=False, data_only=True)
    try:
        products, report = build_products(workbook, feed_products)
    finally:
        workbook.close()
        if temporary_feed:
            feed_path.unlink(missing_ok=True)

    categories = []
    counts = defaultdict(int)
    for product in products:
        counts[product["category"]] += 1
    for order, (name, _glossary, _models) in enumerate(CATEGORY_SHEETS):
        if counts[name]:
            categories.append({"name": name, "order": order, "count": counts[name]})

    payload = {
        "generated": datetime.now().astimezone().isoformat(timespec="seconds"),
        "feed_date": feed_date,
        "source": "FULL_GLOSSARY_USP_TEXTS.xlsx",
        "categories": categories,
        "products": products,
    }
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output / "catalog.json").write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    (args.output / "quality-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"products: {report['models']}")
    print(f"categories: {len(categories)}")
    print(
        f"USP descriptions: {report['described_usps']}/{report['usps']} "
        f"({report['described_usps'] / max(report['usps'], 1):.1%})"
    )
    print(f"feed products: {len(feed_products)}")
    print(f"missing feed models: {len(report['missing_feed_models'])}")
    print(f"missing photos: {len(report['missing_photos'])}")
    print(f"unmatched USP: {len(report['unmatched_usps'])}")
    print(f"ambiguous USP: {len(report['ambiguous_usps'])}")
    print(f"fuzzy matches: {report['fuzzy_matches']}")
    print(f"output: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
