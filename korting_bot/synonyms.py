# -*- coding: utf-8 -*-
from __future__ import annotations

import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from .paths import SYNONYMS_PATH


def norm(s: str) -> str:
    s = (s or "").strip().lower().replace("ё", "е")
    return re.sub(r"\s+", " ", s)


def alnum(s: str) -> str:
    return re.sub(r"[^a-zA-Zа-яА-Я0-9]", "", (s or "").upper().replace("Ё", "Е"))


@dataclass
class Canon:
    id: str
    title: str
    unit: str
    rule: str
    fields: list[str]
    synonyms: list[str]


@dataclass
class Category:
    id: str
    title: str
    prefixes: list[str]
    synonyms: list[str]


@dataclass
class SynonymIndex:
    canons: list[Canon]
    categories: list[Category]
    stop_words: list[str]
    synonym_to_canon: dict[str, Canon] = field(default_factory=dict)
    category_by_syn: dict[str, Category] = field(default_factory=dict)


def _split_pipe(s: str) -> list[str]:
    if not s:
        return []
    return [p.strip() for p in str(s).split("|") if p and str(p).strip()]


def _row_syns(row: tuple, start: int, end: int) -> list[str]:
    out = []
    for i in range(start, end):
        if i >= len(row):
            break
        val = row[i]
        if val is None:
            continue
        t = norm(str(val))
        if t and t not in out:
            out.append(t)
    return out


@lru_cache(maxsize=1)
def load_synonyms(path: str | None = None) -> SynonymIndex:
    xlsx = Path(path) if path else SYNONYMS_PATH
    wb = load_workbook(xlsx, read_only=True, data_only=True)

    canons: list[Canon] = []
    syn_map: dict[str, Canon] = {}
    ws = wb["Каноны"]
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if not row or not row[0]:
            continue
        rule = (row[4] or "value")
        if rule == "skip":
            continue
        fields = _split_pipe(row[5] or "")
        syns = _row_syns(row, 8, 15)
        extra = _split_pipe(row[7] or "")
        for s in extra:
            if s not in syns:
                syns.append(s)
        canon = Canon(
            id=str(row[0]),
            title=str(row[1] or ""),
            unit=str(row[3] or ""),
            rule=str(rule),
            fields=fields,
            synonyms=syns,
        )
        canons.append(canon)
        for s in sorted(syns, key=len, reverse=True):
            key = norm(s)
            if key and key not in syn_map:
                syn_map[key] = canon

    categories: list[Category] = []
    cat_map: dict[str, Category] = {}
    ws = wb["Категории"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        prefixes = [alnum(p) for p in str(row[6] or "").split(",") if p.strip()]
        syns = _row_syns(row, 8, 15)
        cat = Category(id=str(row[0]), title=str(row[1] or ""), prefixes=prefixes, synonyms=syns)
        categories.append(cat)
        for s in syns:
            key = norm(s)
            if key and key not in cat_map:
                cat_map[key] = cat

    stop = []
    if "Стоп-слова" in wb.sheetnames:
        for row in wb["Стоп-слова"].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                stop.append(norm(str(row[0])))
    wb.close()
    return SynonymIndex(
        canons=canons,
        categories=categories,
        stop_words=stop,
        synonym_to_canon=syn_map,
        category_by_syn=cat_map,
    )
